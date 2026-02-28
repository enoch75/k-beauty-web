import sys
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

def markdown_to_docx(md_content, output_path):
    doc = Document()
    
    # Global Font Setting
    style = doc.styles['Normal']
    font = style.font
    font.name = '맑은 고딕'
    font.size = Pt(11)

    lines = md_content.split('\n')
    
    in_table = False
    table_data = []

    for line in lines:
        raw_line = line
        line = line.strip()
        
        # Table detection
        if line.startswith('|') and line.endswith('|'):
            if '---' in line: # Skip separator line
                continue
            cells = [cell.strip() for cell in line.split('|')[1:-1]]
            table_data.append(cells)
            in_table = True
            continue
        else:
            if in_table:
                # Flush table
                if table_data:
                    table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                    table.style = 'Table Grid'
                    for i, row in enumerate(table_data):
                        for j, cell in enumerate(row):
                            table.cell(i, j).text = cell
                table_data = []
                in_table = False

        if not line:
            doc.add_paragraph()
            continue
            
        # Headers
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
        elif line.startswith('#### '):
            doc.add_heading(line[5:], level=3)
            
        # Item lists
        elif re.match(r'^\d+\.', line):
            # Keep original leading spaces for list items if needed, or just add paragraph
            doc.add_paragraph(line) # Simplified for now
            
        # Bullet points
        elif line.startswith('- ') or line.startswith('* '):
            doc.add_paragraph(line[2:], style='List Bullet')
            
        # Normal text
        else:
            doc.add_paragraph(line)
            
    # Handle case where file ends with a table
    if in_table and table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        table.style = 'Table Grid'
        for i, row in enumerate(table_data):
            for j, cell in enumerate(row):
                table.cell(i, j).text = cell

    doc.save(output_path)
    print(f"File successfully saved to: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python md_to_docx.py <input_md_file> <output_docx_file>")
        sys.exit(1)
        
    input_file = sys.argv[1]
    output_file = sys.argv[2]
    
    try:
        with open(input_file, 'r', encoding='utf-8') as f:
            content = f.read()
        markdown_to_docx(content, output_file)
    except Exception as e:
        print(f"Error during conversion: {e}")
        sys.exit(1)
